"""Backend regression tests for StampDoc Romania."""
import io
import os
import time
import uuid
import pytest
import requests
from docx import Document

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://template-stamp-hub.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"


def _build_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    return s


@pytest.fixture(scope="module")
def user_creds():
    uniq = uuid.uuid4().hex[:8]
    return {
        "email": f"test_{uniq}@example.com",
        "password": "TestPass123!",
        "name": "Test User",
        "company": "TestCo",
        "gdpr_consent": True,
    }


@pytest.fixture(scope="module")
def auth(session, user_creds):
    # Register
    r = session.post(f"{API}/auth/register", json=user_creds, timeout=20)
    assert r.status_code == 200, r.text
    data = r.json()
    assert "token" in data and "user" in data
    token = data["token"]
    session.headers.update({"Authorization": f"Bearer {token}"})
    return {"token": token, "user": data["user"]}


# ------------- Health & Plans -------------
class TestRoot:
    def test_root(self, session):
        r = session.get(f"{API}/", timeout=10)
        assert r.status_code == 200
        body = r.json()
        assert body.get("status") == "ok"
        assert body.get("app") == "Energy Project Design Services"
        assert str(body.get("version", "")).startswith("4.5")

    def test_plans(self, session):
        r = session.get(f"{API}/plans", timeout=10)
        assert r.status_code == 200
        plans = r.json()
        # v4.5: plans is a list of 9 public plans in EUR
        assert isinstance(plans, list)
        ids = {p["id"] for p in plans}
        expected = {"basic", "proiectant", "executant", "avize", "ofertare", "contabilitate", "vgd", "rte", "societate"}
        assert expected.issubset(ids), f"missing plans: {expected - ids}"
        assert "developer" not in ids
        by_id = {p["id"]: p for p in plans}
        assert by_id["basic"]["price_eur"] == 99
        assert by_id["proiectant"]["price_eur"] == 149
        assert by_id["societate"]["price_eur"] == 399
        for p in plans:
            assert p["currency"] == "eur"
            for k in ("features", "stamps_allowed", "recipients_allowed", "documents_allowed", "tagline"):
                assert k in p, f"plan {p['id']} missing key {k}"
            assert "export_allowed" in p


# ------------- Auth flow -------------
class TestAuth:
    def test_register_login_me(self, session, auth, user_creds):
        # login
        r = session.post(f"{API}/auth/login", json={
            "email": user_creds["email"], "password": user_creds["password"]
        }, timeout=20)
        assert r.status_code == 200, r.text
        assert "token" in r.json()

        # me
        r = session.get(f"{API}/auth/me", timeout=10)
        assert r.status_code == 200
        u = r.json()
        assert u["email"] == user_creds["email"].lower()
        assert u["plan"] == "basic"
        assert u.get("gdpr_consent") is True
        assert u.get("gdpr_consent_at")

    def test_duplicate_register(self, session, user_creds, auth):
        r = requests.post(f"{API}/auth/register", json=user_creds, timeout=20)
        assert r.status_code == 400
        detail = r.json().get("detail", "")
        assert "deja înregistrat" in detail or "deja" in detail

    def test_unauth_access(self):
        r = requests.get(f"{API}/templates", timeout=10)
        assert r.status_code == 401


# ------------- Templates -------------
class TestTemplates:
    def test_upload_template(self, session, auth):
        docx = _build_docx("Beneficiar: {{nume_beneficiar}}, adresa: {{adresa}}")
        files = {"file": ("tpl.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = session.post(f"{API}/templates/upload", files=files, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["placeholders"] == sorted(["adresa", "nume_beneficiar"])
        pytest.template_id = data["template_id"]

    def test_list_templates(self, session, auth):
        r = session.get(f"{API}/templates", timeout=10)
        assert r.status_code == 200
        ids = [t["template_id"] for t in r.json()]
        assert pytest.template_id in ids

    def test_get_template(self, session, auth):
        r = session.get(f"{API}/templates/{pytest.template_id}", timeout=10)
        assert r.status_code == 200
        assert r.json()["template_id"] == pytest.template_id


# ------------- Stamps -------------
def _make_png() -> bytes:
    from PIL import Image
    img = Image.new("RGBA", (50, 50), (255, 0, 0, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


PNG_BYTES = _make_png()


class TestStamps:
    def test_upload_stamp(self, session, auth):
        files = {"file": ("stamp.png", PNG_BYTES, "image/png")}
        r = session.post(f"{API}/stamps/upload", files=files, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "stamp_id" in data
        pytest.stamp_id = data["stamp_id"]

    def test_list_stamps(self, session, auth):
        r = session.get(f"{API}/stamps", timeout=10)
        assert r.status_code == 200
        ids = [s["stamp_id"] for s in r.json()]
        assert pytest.stamp_id in ids

    def test_get_stamp_image(self, session, auth):
        r = session.get(f"{API}/stamps/{pytest.stamp_id}/image", timeout=10)
        assert r.status_code == 200
        assert r.content[:4] == b"\x89PNG" or len(r.content) > 0


# ------------- Documents Generate -------------
class TestDocuments:
    def test_generate(self, session, auth):
        payload = {
            "template_id": pytest.template_id,
            "values": {"nume_beneficiar": "Ion Popescu", "adresa": "Str. Exemplu 1"},
            "stamp_id": pytest.stamp_id,
            "stamp_position": "bottom-right",
            "stamp_size_cm": 4.0,
            "document_name": "test_doc",
        }
        r = session.post(f"{API}/documents/generate", json=payload, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["stamped"] is True
        pytest.document_id = data["document_id"]

    def test_download(self, session, auth):
        r = session.get(f"{API}/documents/{pytest.document_id}/download", timeout=20)
        assert r.status_code == 200
        # docx is a zip file
        assert r.content[:2] == b"PK"

    def test_list_documents(self, session, auth):
        r = session.get(f"{API}/documents", timeout=10)
        assert r.status_code == 200
        assert any(d["document_id"] == pytest.document_id for d in r.json())


# ------------- Quota -------------
class TestQuota:
    def test_basic_plan_quota_within_limits(self, session, auth):
        # Basic plan = 30 docs/month. Already generated 1 (test_generate). Generate 4 more — all must succeed (well below 30).
        payload = {
            "template_id": pytest.template_id,
            "values": {"nume_beneficiar": "X", "adresa": "Y"},
        }
        for i in range(4):
            r = session.post(f"{API}/documents/generate", json=payload, timeout=30)
            assert r.status_code == 200, f"iter {i}: {r.text}"


# ------------- Email (expect 500: not configured) -------------
class TestEmail:
    def test_email_without_creds(self, session, auth):
        payload = {
            "document_id": pytest.document_id,
            "recipients": ["test@example.com"],
            "subject": "Test",
            "body": "Hello",
        }
        r = session.post(f"{API}/documents/email", json=payload, timeout=20)
        assert r.status_code == 500, r.text
        detail = r.json().get("detail", "")
        # New spec: Romanian 'Setări → Configurare email' substring
        assert "Setări" in detail and "Configurare email" in detail, detail


# ------------- Gmail Config (new) -------------
class TestGmailConfig:
    def test_email_config_initially_false(self, session, auth):
        r = session.get(f"{API}/users/me/email-config", timeout=10)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("configured") is False
        assert body.get("gmail_user") in (None, "")

    def test_patch_users_me_gmail(self, session, auth):
        r = session.patch(f"{API}/users/me", json={
            "gmail_user": "fakeuser@gmail.com",
            "gmail_app_password": "abcdabcdabcdabcd",
        }, timeout=15)
        assert r.status_code == 200, r.text
        user = r.json()
        assert user.get("gmail_configured") is True
        assert user.get("gmail_user") == "fakeuser@gmail.com"
        # NEVER return app password
        assert "gmail_app_password" not in user

    def test_email_config_after_save(self, session, auth):
        r = session.get(f"{API}/users/me/email-config", timeout=10)
        assert r.status_code == 200
        body = r.json()
        assert body["configured"] is True
        assert body["gmail_user"] == "fakeuser@gmail.com"

    def test_me_does_not_return_password(self, session, auth):
        r = session.get(f"{API}/auth/me", timeout=10)
        assert r.status_code == 200
        u = r.json()
        assert "gmail_app_password" not in u
        assert u.get("gmail_configured") is True

    def test_email_with_bad_creds(self, session, auth):
        payload = {
            "document_id": pytest.document_id,
            "recipients": ["test@example.com"],
            "subject": "Test",
            "body": "Hello",
        }
        r = session.post(f"{API}/documents/email", json=payload, timeout=30)
        assert r.status_code == 500, r.text
        detail = r.json().get("detail", "")
        # SMTP auth fails OR generic SMTP error — both acceptable
        assert ("Autentificare Gmail" in detail) or ("SMTP" in detail) or ("Eroare" in detail), detail


# ------------- QES Providers (new) -------------
class TestQESProviders:
    def test_list_providers(self, session, auth):
        r = session.get(f"{API}/qes/providers", timeout=10)
        assert r.status_code == 200, r.text
        providers = r.json()
        assert isinstance(providers, list)
        ids = {p["id"]: p for p in providers}
        assert set(ids.keys()) == {"mock", "certsign", "digisign", "transsped"}, ids.keys()
        assert ids["mock"]["status"] == "active"
        for pid in ("certsign", "digisign", "transsped"):
            assert ids[pid]["status"] == "pending_activation"
            assert isinstance(ids[pid].get("setup_guide"), list)
            assert len(ids[pid]["setup_guide"]) >= 1

    def test_qes_unauth(self):
        r = requests.get(f"{API}/qes/providers", timeout=10)
        assert r.status_code == 401

    def test_set_qes_provider(self, session, auth):
        r = session.patch(f"{API}/users/me", json={"qes_provider": "mock"}, timeout=10)
        assert r.status_code == 200, r.text
        u = r.json()
        assert u.get("qes_provider") == "mock"


# ------------- Stripe Checkout -------------
class TestStripe:
    def test_checkout(self, session, auth):
        payload = {"plan_id": "proiectant", "origin_url": "https://template-stamp-hub.preview.emergentagent.com"}
        r = session.post(f"{API}/payments/checkout", json=payload, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "url" in data and data["url"].startswith("http")
        assert "session_id" in data

    def test_invalid_plan(self, session, auth):
        r = session.post(f"{API}/payments/checkout", json={"plan_id": "bogus", "origin_url": "https://x"}, timeout=15)
        assert r.status_code == 400


# ------------- Delete -------------
class TestDelete:
    def test_delete_document(self, session, auth):
        r = session.delete(f"{API}/documents/{pytest.document_id}", timeout=15)
        assert r.status_code == 200
        # verify gone
        r = session.get(f"{API}/documents/{pytest.document_id}/download", timeout=10)
        assert r.status_code == 404
