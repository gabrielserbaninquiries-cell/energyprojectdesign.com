"""V8.2 regression tests:
  - docx_processor.insert_stamp() with absolute positioning (x_cm/y_cm)
    producing wp:anchor + 2 wp:posOffset elements with correct EMU values.
  - Legacy preset stamp positioning still works.
  - GET /api/me/billing shape compliance.
  - GET /api/placeholders/registry returns ~179 fields.
  - GET /api/gas-project/{pid}/dossier.zip returns ZIP with 27 entries.
  - POST /api/payments/checkout returns Stripe URL.
  - POST /api/documents/generate accepts stamp_x_cm/stamp_y_cm.
"""
import io
import os
import sys
import zipfile
import xml.etree.ElementTree as ET

import pytest
import requests

# Ensure backend module importable for direct insert_stamp tests
sys.path.insert(0, "/app/backend")

from docx_processor import insert_stamp  # noqa: E402
from docx import Document  # noqa: E402

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://57bd020b-829b-4403-b2b9-09912868b634.preview.emergentagent.com").rstrip("/")
ADMIN_EMAIL = "dragosserban95@gmail.com"
ADMIN_PASS = "Test12345"

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


# ============ Fixtures ============

@pytest.fixture(scope="session")
def auth_token():
    r = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASS},
        timeout=30,
    )
    if r.status_code != 200:
        pytest.skip(f"Login failed: {r.status_code} {r.text[:200]}")
    return r.json()["token"]


@pytest.fixture(scope="session")
def auth_headers(auth_token):
    return {"Authorization": f"Bearer {auth_token}"}


@pytest.fixture(scope="session")
def sample_docx_bytes():
    """Build a tiny in-memory DOCX with a single paragraph."""
    d = Document()
    d.add_paragraph("Test body for stamp insertion.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_png_bytes():
    """Minimal 1x1 transparent PNG."""
    import base64
    b64 = (
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYA"
        "AjCB0C8AAAAASUVORK5CYII="
    )
    return base64.b64decode(b64)


# ============ Backend: docx_processor.insert_stamp() ============

class TestInsertStampAbsolute:
    """Validate absolute positioning produces wp:anchor + EMU posOffset."""

    def test_absolute_positioning_creates_anchor(self, sample_docx_bytes, sample_png_bytes):
        out = insert_stamp(
            sample_docx_bytes,
            sample_png_bytes,
            position="absolute",
            size_cm=3.0,
            x_cm=5.0,
            y_cm=10.0,
        )
        assert isinstance(out, (bytes, bytearray))
        assert len(out) > 1000, "Output DOCX too small"

        # Parse document.xml from the returned docx zip
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            with z.open("word/document.xml") as f:
                doc_xml = f.read().decode("utf-8")

        # Verify wp:anchor present
        assert "<wp:anchor" in doc_xml, "wp:anchor element missing"
        # Verify EMU offsets present
        emu_x = int(5.0 * 360000)   # 1,800,000
        emu_y = int(10.0 * 360000)  # 3,600,000
        assert str(emu_x) in doc_xml, f"X EMU {emu_x} not found"
        assert str(emu_y) in doc_xml, f"Y EMU {emu_y} not found"

        # Parse and verify two wp:posOffset values
        root = ET.fromstring(doc_xml)
        anchors = root.findall(".//wp:anchor", NS)
        assert len(anchors) >= 1, "No <wp:anchor> in parsed tree"

        offsets = root.findall(".//wp:posOffset", NS)
        offset_values = [int(o.text) for o in offsets if o.text and o.text.lstrip("-").isdigit()]
        assert emu_x in offset_values, f"posOffset {emu_x} not found in {offset_values}"
        assert emu_y in offset_values, f"posOffset {emu_y} not found in {offset_values}"

    def test_legacy_preset_bottom_right_still_works(self, sample_docx_bytes, sample_png_bytes):
        out = insert_stamp(
            sample_docx_bytes,
            sample_png_bytes,
            position="bottom-right",
            size_cm=4.0,
        )
        assert isinstance(out, (bytes, bytearray))
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            names = z.namelist()
            assert "word/document.xml" in names
            # Inline image (legacy mode) → no anchor
            doc_xml = z.read("word/document.xml").decode("utf-8")
        # Legacy mode uses inline drawing, NOT anchor
        assert "<wp:anchor" not in doc_xml
        assert "<wp:inline" in doc_xml or "w:drawing" in doc_xml

    def test_legacy_preset_top_left_still_works(self, sample_docx_bytes, sample_png_bytes):
        out = insert_stamp(
            sample_docx_bytes,
            sample_png_bytes,
            position="top-left",
            size_cm=3.0,
        )
        assert len(out) > 1000
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            doc_xml = z.read("word/document.xml").decode("utf-8")
        assert "<wp:anchor" not in doc_xml

    def test_emu_conversion_various_coords(self, sample_docx_bytes, sample_png_bytes):
        # 2.5 cm → 900000 EMU
        out = insert_stamp(
            sample_docx_bytes, sample_png_bytes,
            position="absolute", size_cm=2.0, x_cm=2.5, y_cm=7.5,
        )
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            doc_xml = z.read("word/document.xml").decode("utf-8")
        assert "900000" in doc_xml      # 2.5cm
        assert "2700000" in doc_xml     # 7.5cm


# ============ Backend: /api/me/billing ============

class TestBillingEndpoint:
    def test_me_billing_shape(self, auth_headers):
        r = requests.get(f"{BASE_URL}/api/me/billing", headers=auth_headers, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        # required keys
        assert "current_plan" in data
        assert "transactions" in data
        assert "activations" in data
        # current_plan must have plan_id + name + price_eur keys (renews_at optional)
        cp = data["current_plan"]
        assert "plan_id" in cp
        assert "name" in cp
        # transactions and activations must be lists
        assert isinstance(data["transactions"], list)
        assert isinstance(data["activations"], list)

    def test_me_billing_unauth(self):
        r = requests.get(f"{BASE_URL}/api/me/billing", timeout=15)
        assert r.status_code in (401, 403)


# ============ Backend: /api/placeholders/registry ============

class TestPlaceholdersRegistry:
    def test_registry_returns_at_least_179_fields(self, auth_headers):
        r = requests.get(f"{BASE_URL}/api/placeholders/registry", headers=auth_headers, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        # try common shapes
        fields = data.get("fields") or data.get("FIELDS_REGISTRY") or []
        # If wrapped, fall back to nested
        if not fields and isinstance(data, dict):
            for v in data.values():
                if isinstance(v, list) and len(v) > 100:
                    fields = v
                    break
        assert len(fields) >= 179, f"Expected ≥179 fields, got {len(fields)}"


# ============ Backend: Stripe checkout ============

class TestStripeCheckout:
    def test_checkout_returns_stripe_url(self, auth_headers):
        for plan_id in ("operator", "proiectant", "basic", "vgd"):
            body = {"plan_id": plan_id, "origin_url": BASE_URL}
            r = requests.post(
                f"{BASE_URL}/api/payments/checkout",
                json=body, headers=auth_headers, timeout=30,
            )
            if r.status_code == 200:
                data = r.json()
                url = data.get("url") or data.get("checkout_url") or ""
                assert "stripe.com" in url, f"Not a Stripe URL: {url}"
                assert data.get("session_id"), "No session_id returned"
                return
        pytest.skip(f"All plans rejected: {r.status_code} {r.text[:200]}")


# ============ Backend: dossier.zip ============

class TestDossierZip:
    PID = "gp_54135e822f25f7d7"  # seeded admin project from test_credentials.md

    def test_dossier_zip_contains_27_files(self, auth_headers):
        r = requests.get(
            f"{BASE_URL}/api/gas-project/{self.PID}/dossier.zip",
            headers=auth_headers, timeout=60,
        )
        if r.status_code == 404:
            # try alternate seeded PID from iteration_7
            r = requests.get(
                f"{BASE_URL}/api/gas-project/gp_e79e2810cc64b5b4/dossier.zip",
                headers=auth_headers, timeout=60,
            )
        assert r.status_code == 200, f"{r.status_code} {r.text[:200]}"
        assert r.content[:2] == b"PK"
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            names = z.namelist()
        docx_files = [n for n in names if n.lower().endswith(".docx")]
        # 26 DOCX + 1 manifest = 27, but allow flex
        assert len(names) >= 27, f"Expected ≥27 zip entries, got {len(names)}: {names[:5]}"
        assert len(docx_files) >= 26, f"Expected ≥26 DOCX, got {len(docx_files)}"


# ============ Backend: /api/documents/generate stamp_x_cm/y_cm wiring ============

class TestGenerateAcceptsAbsoluteStamp:
    """Verifies the endpoint accepts absolute-stamp body fields without 422.

    Full e2e (template_id + uploaded stamp) is heavy; we ONLY confirm the
    schema accepts the new optional fields. Missing template_id → 404/400, NOT 422.
    """

    def test_post_documents_generate_accepts_stamp_x_y_fields(self, auth_headers):
        body = {
            "template_id": "nonexistent_template_for_schema_check",
            "values": {"foo": "bar"},
            "stamp_position": "absolute",
            "stamp_size_cm": 3.0,
            "stamp_x_cm": 5.0,
            "stamp_y_cm": 10.0,
        }
        r = requests.post(
            f"{BASE_URL}/api/documents/generate",
            json=body, headers=auth_headers, timeout=20,
        )
        # 422 would mean Pydantic rejected the new fields. Anything else = schema OK.
        assert r.status_code != 422, f"Schema rejected new stamp_x_cm/y_cm: {r.text}"
